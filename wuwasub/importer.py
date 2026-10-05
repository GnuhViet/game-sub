"""Đọc bộ sub EN–VI từ CSV/TSV/TXT/XLSX/JSON thành bảng (headers, rows) và đoán cột."""
import csv, json, io, re
from pathlib import Path

VI_CHARS = re.compile(r"[ăâđêôơưạảấầẩẫậắằẳẵặẹẻẽếềểễệỉịọỏốồổỗộớờởỡợụủứừửữựỳỵỷỹáàãéèíìóòõúùýĂÂĐÊÔƠƯ]")
SRC_NAMES = ("en", "eng", "english", "source", "src", "original", "goc", "gốc", "text_en", "en_us", "content")
VI_NAMES = ("vi", "vn", "viet", "vietnamese", "tiếng việt", "dich", "dịch", "translation", "target", "text_vi", "vi_vn")

def _decode(raw: bytes) -> str:
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"): return raw.decode("utf-16")
    for enc in ("utf-8-sig", "cp1258", "latin-1"):
        try: return raw.decode(enc)
        except UnicodeDecodeError: pass

def _flatten(obj, prefix=""):
    if isinstance(obj, dict):
        out = {}
        for k, v in obj.items():
            if isinstance(v, (dict, list)): out.update(_flatten(v, f"{prefix}{k}."))
            else: out[f"{prefix}{k}"] = v
        return out
    return {prefix.rstrip(".") or "value": obj}

def read_table(path, sheet=None):
    p = Path(path); ext = p.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        ws = wb[sheet] if sheet else next((w for w in wb.worksheets if w.max_row and w.max_row > 1), wb.worksheets[0])
        rows = [["" if c is None else str(c) for c in r] for r in ws.iter_rows(values_only=True)]
        wb.close()
        return _with_header([r for r in rows if any(x.strip() for x in r)])
    text = _decode(p.read_bytes())
    if ext == ".json":
        data = json.loads(text)
        if isinstance(data, dict):
            vals = list(data.values())
            if vals and all(isinstance(v, str) for v in vals): return ["key", "value"], [[k, v] for k, v in data.items()]
            data = [dict(_flatten(v), _id=k) if isinstance(v, dict) else {"_id": k, "value": v} for k, v in data.items()]
        if data and isinstance(data[0], list): return _with_header([[str(x) for x in r] for r in data])
        recs = [_flatten(r) for r in data]; hdr = list(dict.fromkeys(k for r in recs for k in r))
        return hdr, [["" if r.get(h) is None else str(r.get(h)) for h in hdr] for r in recs]
    # csv / tsv / txt
    sample = text[:20000]
    if ext == ".tsv" or "\t" in sample: delim = "\t"
    else:
        try: delim = csv.Sniffer().sniff(sample, delimiters=",;|").delimiter
        except csv.Error: delim = ","
    rows = [r for r in csv.reader(io.StringIO(text), delimiter=delim) if any(x.strip() for x in r)]
    return _with_header(rows)

def _with_header(rows):
    if not rows: return [], []
    first = [x.strip().lower() for x in rows[0]]
    if any(n in first for n in SRC_NAMES + VI_NAMES) or (len(rows) > 1 and not any(VI_CHARS.search(x) for x in rows[0]) and all(len(x) < 30 for x in rows[0])):
        hdr, body = [x.strip() or f"col{i+1}" for i, x in enumerate(rows[0])], rows[1:]
    else: hdr, body = [f"col{i+1}" for i in range(max(map(len, rows)))], rows
    w = len(hdr)
    return hdr, [(r + [""] * w)[:w] for r in body]

def guess_cols(hdr, rows):
    """Trả về (src_idx, vi_idx)."""
    low = [h.lower().split(".")[-1] for h in hdr]
    src = next((i for i, h in enumerate(low) if h in SRC_NAMES), None)
    vi = next((i for i, h in enumerate(low) if h in VI_NAMES), None)
    if src is not None and vi is not None and src != vi: return src, vi
    sample = rows[:300]; n = len(hdr)
    vi_ratio = [sum(bool(VI_CHARS.search(r[i])) for r in sample) / max(1, len(sample)) for i in range(n)]
    ascii_ratio = [sum(bool(r[i].strip()) and r[i].isascii() and bool(re.search(r"[A-Za-z]{2}", r[i])) for r in sample) / max(1, len(sample)) for i in range(n)]
    if vi is None: vi = max(range(n), key=lambda i: vi_ratio[i]) if n else 0
    if src is None: src = max((i for i in range(n) if i != vi), key=lambda i: ascii_ratio[i], default=0)
    return src, vi

def extract_pairs(rows, src_idx, vi_idx):
    out = []
    for r in rows:
        s, v = r[src_idx].strip(), r[vi_idx].strip()
        if s and v: out.append((s, v))
    return out
